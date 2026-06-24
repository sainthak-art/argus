"""표준 사전 파일(CSV/XLSX) 파서.

헤더 행을 키로 하는 행 dict 목록을 반환한다. 값은 모두 trim 된 문자열.
빈 행은 제외한다. 컬럼명→스키마 필드 매핑은 service 레이어가 담당.
"""

from __future__ import annotations

import csv
import io


def parse_table(filename: str, content: bytes) -> list[dict[str, str]]:
    """파일 확장자로 CSV/XLSX 를 판별해 행 dict 목록을 반환."""
    name = (filename or "").lower()
    if name.endswith(".csv"):
        return _parse_csv(content)
    return _parse_xlsx(content)


def _parse_csv(content: bytes) -> list[dict[str, str]]:
    text = content.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    rows: list[dict[str, str]] = []
    for raw in reader:
        row = {
            (k or "").strip(): ("" if v is None else str(v).strip())
            for k, v in raw.items()
            if k
        }
        if any(row.values()):
            rows.append(row)
    return rows


def _parse_xlsx(content: bytes) -> list[dict[str, str]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        try:
            header_cells = next(rows_iter)
        except StopIteration:
            return []
        headers = ["" if h is None else str(h).strip() for h in header_cells]

        out: list[dict[str, str]] = []
        for values in rows_iter:
            if values is None:
                continue
            cells = ["" if v is None else str(v).strip() for v in values]
            if not any(cells):
                continue
            row = {
                headers[i]: (cells[i] if i < len(cells) else "")
                for i in range(len(headers))
                if headers[i]
            }
            out.append(row)
        return out
    finally:
        wb.close()
